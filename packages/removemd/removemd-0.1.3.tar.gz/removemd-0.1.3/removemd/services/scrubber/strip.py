import datetime, tempfile, os, subprocess
from PIL import Image
from PyPDF2 import PdfReader, PdfWriter
from io import BytesIO
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

def strip_image(file, fmt):
    try:
        image = Image.open(file)
        cleaned_image_io = BytesIO()
        if fmt.upper() == "JPEG":
            image.save(cleaned_image_io, format=fmt, quality=95, optimize=True)
        else:
            image.save(cleaned_image_io, format=fmt)
        cleaned_image_io.seek(0)
        return cleaned_image_io
    except Exception as e:
        return f"Error processing image file: {str(e)}"

def strip_pdf(file):
    try:
        reader = PdfReader(file)
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        cleaned_pdf = BytesIO()
        writer.write(cleaned_pdf)
        cleaned_pdf.seek(0)
        return cleaned_pdf
    except Exception as e:
        return f"Error processing PDF file: {str(e)}"

def strip_office(file, fmt):
    try:
        cleaned_office = BytesIO()
        file.seek(0)
        text_attrs = [
            "author", "title", "subject", "keywords", "comments", "last_modified_by",
            "category", "content_status", "identifier", "language", "version", "creator",
            "description", "manager", "company", "hyperlink_base"
        ]
        now = datetime.datetime.now()

        if fmt == "DOCX":
            doc = Document(file)
            for attr in text_attrs:
                setattr(doc.core_properties, attr, "")
            doc.core_properties.revision = 1
            doc.core_properties.created = now
            doc.core_properties.modified = now
            doc.save(cleaned_office)

        elif fmt == "XLSX":
            wb = load_workbook(file)
            for attr in text_attrs:
                setattr(wb.properties, attr, "")
            wb.properties.revision = 1
            wb.properties.created = now
            wb.properties.modified = now
            wb.save(cleaned_office)

        elif fmt == "PPTX":
            prs = Presentation(file)
            for attr in text_attrs:
                setattr(prs.core_properties, attr, "")
            prs.core_properties.revision = 1
            prs.core_properties.created = now
            prs.core_properties.modified = now
            prs.save(cleaned_office)

        else:
            return f"Unsupported Office file format: {fmt}"

        cleaned_office.seek(0)
        return cleaned_office
    except Exception as e:
        return f"Error processing Office file: {str(e)}"

def strip_audio(file, fmt):
    supported_formats = [
        "mp3", "flac", "ogg", "opus", "wav", "aac", "m4a", 
        "wma", "aiff", "alac", "ape", "mpc", "wv"
    ]
    if fmt.lower() not in supported_formats:
        return "Format not supported."
    
    file.seek(0)
    file_content = file.read()
    with tempfile.NamedTemporaryFile(delete=False, suffix=f'.{fmt}') as temp_input:
        temp_input.write(file_content)
        temp_input_name = temp_input.name
    with tempfile.NamedTemporaryFile(delete=False, suffix=f'.{fmt}') as temp_output:
        temp_output_name = temp_output.name
    
    try:
        cmd = [
            'ffmpeg',
            '-i', temp_input_name,
            '-map_metadata', '-1',
            '-c', 'copy',
            '-loglevel', 'error',
            '-y',
            temp_output_name
        ]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=30)
        if result.returncode != 0:
            return f"Error processing audio file with FFmpeg: {result.stderr}"
        if not os.path.exists(temp_output_name) or os.path.getsize(temp_output_name) == 0:
            return "Error processing audio file - empty output"
        cleaned_audio = BytesIO()
        with open(temp_output_name, 'rb') as f:
            cleaned_audio.write(f.read())
        cleaned_audio.seek(0)
        return cleaned_audio
    except subprocess.TimeoutExpired:
        return "Audio processing timed out"
    except Exception as e:
        return f"Error during audio processing: {str(e)}"
    finally:
        for temp_file in [temp_input_name, temp_output_name]:
            try:
                if os.path.exists(temp_file):
                    os.unlink(temp_file)
            except:
                pass

def strip_video(file, fmt):
    supported_formats = [
        "mp4", "mkv", "avi", "mov", "wmv", "flv", "webm", "mpeg", "mpg", "3gp"
    ]
    if fmt.lower() not in supported_formats:
        return "Format not available for use for the moment."
    
    file.seek(0)
    with tempfile.NamedTemporaryFile(delete=False, suffix=f'.{fmt}') as temp_input:
        temp_input.write(file.read())
        temp_input_name = temp_input.name
    with tempfile.NamedTemporaryFile(delete=False, suffix=f'.{fmt}') as temp_output:
        temp_output_name = temp_output.name
    
    try:
        cmd = [
            'ffmpeg',
            '-i', temp_input_name,
            '-map', '0',
            '-c', 'copy',
            '-movflags', 'faststart',
            '-map_metadata', '-1',
            '-loglevel', 'error',
            '-y',
            temp_output_name
        ]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=30)
        if result.returncode != 0:
            return f"Error processing video file: {result.stderr}"
        cleaned_video = BytesIO()
        with open(temp_output_name, 'rb') as f:
            cleaned_video.write(f.read())
        cleaned_video.seek(0)
        return cleaned_video
    except subprocess.TimeoutExpired:
        return "Error processing video file, operation timed out."
    except Exception as e:
        return f"Error processing video file: {str(e)}"
    finally:
        for temp_file in [temp_input_name, temp_output_name]:
            try:
                if os.path.exists(temp_file):
                    os.unlink(temp_file)
            except:
                pass
