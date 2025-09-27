from PIL import Image, ExifTags
from PyPDF2 import PdfReader
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from mutagen import File as MutagenFile
import tempfile, subprocess, os, json
from PIL import Image, ExifTags

def check_image(file):
    try:
        img = Image.open(file)
        meta = {}

        # General image info
        if img.info:
            clean_info = {
                k: str(v) for k, v in img.info.items()
                if isinstance(v, (str, int, float)) and not str(v).startswith("b'")
            }
            if clean_info:
                meta["General Info"] = clean_info

        # EXIF metadata
        exif_data = getattr(img, "_getexif", lambda: None)()
        if exif_data:
            exif = {}
            for tag_id, value in exif_data.items():
                tag = ExifTags.TAGS.get(tag_id, tag_id)
                val_str = str(value)
                if not val_str.startswith("b'") and not val_str.startswith("Exif"):
                    exif[tag] = val_str
            if exif:
                meta["EXIF Data"] = exif

        # Final output
        return {"metadata": meta if meta else "No readable metadata found"}

    except Exception:
        return "Error checking image metadata"

def check_pdf(file):
    try:
        reader = PdfReader(file)
        meta = {}
        if reader.metadata:
            meta = {k: str(v) for k, v in reader.metadata.items()}
        return {"metadata": meta if meta else "No metadata found"}
    except Exception:
        return "Error checking PDF metadata"

def check_office(file, fmt):
    try:
        file.seek(0)
        meta = {}
        if fmt == "DOCX":
            core = Document(file).core_properties
        elif fmt == "XLSX":
            core = load_workbook(file).properties
        elif fmt == "PPTX":
            core = Presentation(file).core_properties
        else:
            return f"Unsupported Office format {fmt}"
        for attr in dir(core):
            if not attr.startswith("_"):
                try:
                    value = getattr(core, attr)
                    if value not in [None, ""]:
                        meta[attr] = str(value)
                except:
                    pass
        return {"metadata": meta if meta else "No metadata found"}
    except Exception:
        return "Error checking Office metadata"

def check_audio(file):
    try:
        file.seek(0)
        media = MutagenFile(file)
        meta = {}
        if media and media.tags:
            for k, v in media.tags.items():
                meta[k] = str(v)
        return {"metadata": meta if meta else "No metadata found"}
    except Exception:
        return "Error checking audio metadata"

def check_video(file, fmt):
    try:
        file.seek(0)
        with tempfile.NamedTemporaryFile(delete=False, suffix=f".{fmt}") as temp:
            temp.write(file.read())
            temp_name = temp.name

        cmd = [
            "ffprobe", "-v", "quiet", "-print_format", "json",
            "-show_format", "-show_streams", temp_name
        ]
        result = subprocess.run(cmd, capture_output=True, text=True)
        os.unlink(temp_name)

        if result.returncode == 0:
            data = json.loads(result.stdout)
            meta = {
                "filename": data["format"]["filename"],
                "duration": float(data["format"]["duration"]),
                "size": int(data["format"]["size"]),
                "bit_rate": int(data["format"]["bit_rate"]),
                "streams": []
            }
            for s in data["streams"]:
                stream_meta = {
                    "index": s.get("index"),
                    "codec_type": s.get("codec_type"),
                    "codec_name": s.get("codec_name"),
                    "width": s.get("width"),
                    "height": s.get("height"),
                    "bit_rate": int(s["bit_rate"]) if s.get("bit_rate") else None,
                    "duration": float(s["duration"]) if s.get("duration") else None,
                    "r_frame_rate": s.get("r_frame_rate"),
                }
                meta["streams"].append({k:v for k,v in stream_meta.items() if v is not None})
            return {"metadata": meta}
        else:
            return "Error extracting video metadata"
    except Exception as e:
        return f"Error checking video metadata: {str(e)}"

